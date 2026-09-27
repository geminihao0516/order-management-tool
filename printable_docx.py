"""把展開後的訂單製成 A4 雙欄 Word 列印檔。"""

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def _set_table_borders(table):
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '12')
        border.set(qn('w:color'), '000000')
        borders.append(border)
    table._tbl.tblPr.append(borders)


def _fill_cell(cell, order):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    lines = [f"{order['index']}　{order['item']}　{order['main_person']}"]
    target = str(order.get('target_person', '')).strip()
    wish = str(order.get('wish', '')).strip()
    if target and target not in ('—', '-'):
        lines.append(target)
    if wish:
        lines.append(wish)

    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run('\n'.join(lines))
    run.font.name = 'Microsoft JhengHei'
    run.font.size = Pt(11)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')


def build_printable_docx(expanded_orders):
    """回傳與附件相近的 A4、兩欄、逐支蠟燭列印版 Word bytes。"""
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.27)
    section.bottom_margin = Cm(1.27)
    section.left_margin = Cm(1.27)
    section.right_margin = Cm(1.27)

    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column in table.columns:
        column.width = Cm(9.23)
    _set_table_borders(table)

    for offset in range(0, len(expanded_orders), 2):
        row = table.add_row()
        for cell in row.cells:
            cell.width = Cm(9.23)
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        _fill_cell(row.cells[0], expanded_orders[offset])
        if offset + 1 < len(expanded_orders):
            _fill_cell(row.cells[1], expanded_orders[offset + 1])

    output = BytesIO()
    document.save(output)
    return output.getvalue()
