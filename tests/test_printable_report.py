import unittest
from io import BytesIO
from zipfile import ZipFile

import fitz
from docx import Document

from order_formatter import OrderFormatter
from printable_docx import build_printable_docx
from printable_pdf import build_a4_png_zip, build_printable_pdf


class PrintableReportTests(unittest.TestCase):
    def test_numbered_input_keeps_wishes_with_their_item_group(self):
        data = '''1. 測試甲 1988/7/15
九尾狐1、粉色燕通1
願望：第一行
第二行

2. 測試乙 1984/1/23
🕯️象神1
願望：工作
🕯️巴拉吉1
願望：銷售
🕯️九尾狐1、粉色燕通1
願望：感情
'''

        formatter = OrderFormatter()
        formatter.load_data(data)

        self.assertEqual(formatter.customer_count, 2)
        self.assertEqual(len(formatter.orders), 4)
        self.assertEqual(len(formatter.expanded_orders), 6)
        self.assertEqual(
            [order['wish'] for order in formatter.orders],
            ['第一行 第二行', '工作', '銷售', '感情'],
        )
        self.assertEqual(sum(formatter.item_amounts.values()), 1900)
        self.assertEqual(formatter.PRICE_LIST['象神'], 350)

    def test_word_output_has_two_cells_per_row(self):
        formatter = OrderFormatter()
        formatter.load_data('九尾狐x3\t測試甲/1988.07.15\t—\t測試願望')

        document = Document(BytesIO(build_printable_docx(formatter.expanded_orders)))
        table = document.tables[0]

        self.assertEqual(len(table.columns), 2)
        self.assertEqual(len(table.rows), 2)
        self.assertTrue(table.cell(0, 0).text.startswith('1\u3000九尾狐'))
        self.assertTrue(table.cell(1, 0).text.startswith('3\u3000九尾狐'))
        self.assertEqual(table.cell(1, 1).text, '')

    def test_pdf_and_png_are_a4_and_keep_every_item(self):
        formatter = OrderFormatter()
        formatter.load_data('九尾狐x31\t測試甲/1988.07.15\t—\t' + '測試願望' * 40)

        pdf_bytes = build_printable_pdf(formatter.expanded_orders)
        with fitz.open(stream=pdf_bytes, filetype='pdf') as document:
            self.assertGreater(len(document), 1)
            self.assertEqual((document[0].rect.width, document[0].rect.height), (595, 842))
            content = ''.join(page.get_text() for page in document)
            self.assertIn('九尾狐', content)
            self.assertIn('測試甲/1988.07.15', content)
            self.assertIn('31', content)
            page_count = len(document)

        with ZipFile(BytesIO(build_a4_png_zip(pdf_bytes))) as archive:
            self.assertEqual(len(archive.namelist()), page_count)
            first_image = fitz.Pixmap(archive.read(archive.namelist()[0]))
            self.assertGreaterEqual(first_image.width, 2479)
            self.assertGreaterEqual(first_image.height, 3508)


if __name__ == '__main__':
    unittest.main()
